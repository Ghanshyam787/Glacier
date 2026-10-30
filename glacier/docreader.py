"""
Glacier - Document reader
PDF, Word (.docx), Excel (.xlsx), CSV, HTML aur saari text/code files se text nikalta hai.
Optional libraries (pypdf, python-docx, openpyxl) na ho to us format ke liye saaf error deta hai.
"""

import csv
import io
import json
import os

import netsafe

TEXT_EXTS = {
    ".txt", ".md", ".markdown", ".rst", ".log", ".ini", ".cfg", ".conf", ".toml", ".yaml", ".yml", ".json",
    ".xml", ".csv", ".tsv", ".py", ".js", ".ts", ".jsx", ".tsx", ".java", ".c", ".cpp", ".h", ".hpp", ".cs",
    ".go", ".rs", ".rb", ".php", ".sh", ".bat", ".ps1", ".sql", ".css", ".scss", ".kt", ".swift", ".r",
    ".tex", ".srt", ".vtt", ".env.example",
}
HTML_EXTS = {".html", ".htm"}
INDEXABLE_EXTS = TEXT_EXTS | HTML_EXTS | {".pdf", ".docx", ".xlsx"}
MAX_FILE_BYTES = 40 * 1024 * 1024


def _read_plain(path: str) -> str:
    with open(path, "rb") as f:
        raw = f.read(8_000_000)
    if raw.startswith((b"\xff\xfe", b"\xfe\xff")):
        return raw.decode("utf-16", errors="replace")
    for enc in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="replace")


def read_text(path: str, max_chars: int = 20000):
    """Returns (text, error). Exactly one of them is None."""
    if not os.path.isfile(path):
        return None, f"File nahi mili: {path}"
    try:
        if os.path.getsize(path) > MAX_FILE_BYTES:
            return None, "File bahut badi hai (40MB se zyada)."
    except OSError as e:
        return None, f"File access error: {e}"

    ext = os.path.splitext(path)[1].lower()
    try:
        if ext == ".pdf":
            try:
                from pypdf import PdfReader
            except ImportError:
                return None, "PDF padhne ke liye 'pypdf' chahiye (pip install pypdf)."
            reader = PdfReader(path)
            if reader.is_encrypted:
                try:
                    reader.decrypt("")
                except Exception:
                    return None, "PDF password-protected hai."
            parts = []
            for i, page in enumerate(reader.pages[:80]):
                parts.append(page.extract_text() or "")
                if sum(len(p) for p in parts) > max_chars * 2:
                    break
            text = "\n".join(parts).strip()
            if not text:
                return None, "PDF mein text nahi mila (scanned image lagti hai - look_at_image se padho)."
            return text[:max_chars], None

        if ext == ".docx":
            try:
                import docx
            except ImportError:
                return None, "Word file ke liye 'python-docx' chahiye (pip install python-docx)."
            d = docx.Document(path)
            parts = [p.text for p in d.paragraphs if p.text.strip()]
            for table in d.tables[:10]:
                for row in table.rows:
                    parts.append(" | ".join(c.text.strip() for c in row.cells))
            return "\n".join(parts)[:max_chars], None

        if ext == ".xlsx":
            try:
                import openpyxl
            except ImportError:
                return None, "Excel ke liye 'openpyxl' chahiye (pip install openpyxl)."
            wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
            lines = []
            for ws in wb.worksheets[:5]:
                lines.append(f"## Sheet: {ws.title}")
                for i, row in enumerate(ws.iter_rows(values_only=True)):
                    if i >= 200:
                        lines.append("…")
                        break
                    lines.append(" | ".join("" if c is None else str(c) for c in row))
            return "\n".join(lines)[:max_chars], None

        if ext in HTML_EXTS:
            return netsafe.html_to_text(_read_plain(path), max_chars), None

        if ext in (".csv", ".tsv"):
            text = _read_plain(path)
            delim = "\t" if ext == ".tsv" else ","
            rows = list(csv.reader(io.StringIO(text), delimiter=delim))
            return "\n".join(" | ".join(r) for r in rows[:300])[:max_chars], None

        if ext == ".json":
            text = _read_plain(path)
            try:
                text = json.dumps(json.loads(text), ensure_ascii=False, indent=1)
            except Exception:
                pass
            return text[:max_chars], None

        if ext in TEXT_EXTS or ext == "":
            return _read_plain(path)[:max_chars], None

        return None, f"'{ext}' format text ke roop mein nahi padh sakta."
    except Exception as e:
        return None, f"File padhne mein error: {type(e).__name__}: {str(e)[:200]}"
